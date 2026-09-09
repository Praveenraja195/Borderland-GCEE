"""
Bulk team import — the parsing, detection and derivation rules.

Everything covered here is a pure function of its input, so this module needs
no Postgres and no Redis (unlike the rest of the suite, which is integration-
first against a real database because the schema's triggers are half the
business logic). What it does need to pin down is the stuff that silently
ruins an event if it's wrong: a password derived from the country code
instead of the number, a column mapped to the wrong field, a duplicate
slipping past into a UNIQUE violation halfway through the import.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest
from openpyxl import Workbook, load_workbook

from app.services import export_service, team_import_service as tis


# ---------------------------------------------------------------------------
# Phone normalisation and the derived password
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("9876543210", "9876543210"),
        ("98765 43210", "9876543210"),
        ("+91 98765 43210", "9876543210"),      # country code dropped
        ("+919876543210", "9876543210"),
        ("091-98765-43210", "9876543210"),      # trunk prefix dropped
        ("(987) 654-3210", "9876543210"),
        (9876543210, "9876543210"),             # Excel numeric cell
        (9876543210.0, "9876543210"),           # Excel float cell
        ("", ""),
        (None, ""),
        ("not a phone", ""),
    ],
)
def test_normalise_phone(raw, expected):
    assert tis.normalise_phone(raw) == expected


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("9876543210", "9876"),
        ("+91 98765 43210", "9876"),   # NOT "9198" — the whole point of trimming
        ("+919123456780", "9123"),
        (9876543210.0, "9876"),
        ("123", ""),                   # too short to derive from
        ("", ""),
        (None, ""),
    ],
)
def test_password_from_phone(raw, expected):
    assert tis.password_from_phone(raw) == expected


def test_password_is_first_four_digits_not_first_four_characters():
    """A number written with a leading zero or spacing must not leak
    formatting into the password."""
    assert tis.password_from_phone(" 98-76 54 3210 ") == "9876"


# ---------------------------------------------------------------------------
# Team code format
# ---------------------------------------------------------------------------


def test_generate_team_code_format():
    code = tis.generate_team_code(set())
    assert code.startswith("B@GCEE-")
    assert code.endswith("#")
    digits = code[len("B@GCEE-") : -1]
    assert len(digits) == 4 and digits.isdigit()


def test_generate_team_code_is_unique_within_a_run():
    taken: set[str] = set()
    codes = {tis.generate_team_code(taken) for _ in range(200)}
    assert len(codes) == 200
    assert codes == taken


def test_generate_team_code_avoids_existing_codes():
    taken = {f"B@GCEE-{n:04d}#" for n in range(0, 9999)}
    remaining = "B@GCEE-9999#"
    assert tis.generate_team_code(taken) == remaining


# ---------------------------------------------------------------------------
# Column detection
# ---------------------------------------------------------------------------


def test_detect_columns_plain_headers():
    mapping = tis.detect_columns(["Team Name", "Leader Name", "Phone Number", "Email"])
    assert mapping == {"team_name": 0, "leader_name": 1, "leader_phone": 2, "leader_email": 3}


def test_detect_columns_messy_google_form_headers():
    headers = [
        "Timestamp",
        "Email Address",
        "Name of the Team",
        "Team Leader's Full Name",
        "WhatsApp Number",
        "Number of members",
    ]
    mapping = tis.detect_columns(headers)
    assert mapping["team_name"] == 2
    assert mapping["leader_name"] == 3
    assert mapping["leader_phone"] == 4
    assert mapping["leader_email"] == 1


def test_detect_columns_does_not_confuse_team_name_with_leader_name():
    """'Team Name' and 'Team Leader Name' both contain 'team' and 'name' —
    the more specific synonym has to win each column."""
    mapping = tis.detect_columns(["Team Leader Name", "Team Name"])
    assert mapping["leader_name"] == 0
    assert mapping["team_name"] == 1


def test_detect_columns_one_column_is_never_claimed_twice():
    mapping = tis.detect_columns(["Team", "Leader", "Mobile", "Mail"])
    assert len(set(mapping.values())) == len(mapping)


def test_apply_overrides_by_name_letter_and_index():
    headers = ["A col", "B col", "C col", "D col"]
    base = {"team_name": 0}
    assert tis.apply_overrides(base, headers, {"leader_phone": "C col"})["leader_phone"] == 2
    assert tis.apply_overrides(base, headers, {"leader_phone": "C"})["leader_phone"] == 2
    assert tis.apply_overrides(base, headers, {"leader_phone": "2"})["leader_phone"] == 2


def test_apply_overrides_rejects_unknown_column():
    from app.core.exceptions import AppError

    with pytest.raises(AppError):
        tis.apply_overrides({}, ["A", "B"], {"leader_phone": "Nope"})


# ---------------------------------------------------------------------------
# Sheet parsing
# ---------------------------------------------------------------------------


def _xlsx_bytes(rows) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registrations"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parse_sheet_xlsx():
    data = _xlsx_bytes(
        [
            ["Team Name", "Leader Name", "Phone Number", "Email"],
            ["Dragon Warriors", "Praveen", "9876543210", "p@example.com"],
        ]
    )
    parsed = tis.parse_sheet(data, "teams.xlsx")
    assert parsed.headers == ["Team Name", "Leader Name", "Phone Number", "Email"]
    assert parsed.header_row_number == 1
    assert len(parsed.rows) == 1
    assert parsed.sheet_name == "Registrations"


def test_parse_sheet_skips_a_title_banner_above_the_table():
    """Registration sheets routinely open with a merged title row. The header
    row is found by content, not assumed to be row 1."""
    data = _xlsx_bytes(
        [
            ["GCEE 2026 — Team Registrations", None, None, None],
            [None, None, None, None],
            ["Team Name", "Leader Name", "Phone Number", "Email"],
            ["Dragon Warriors", "Praveen", "9876543210", "p@example.com"],
        ]
    )
    parsed = tis.parse_sheet(data, "teams.xlsx")
    assert parsed.header_row_number == 3
    assert parsed.headers[0] == "Team Name"
    assert len(parsed.rows) == 1


def test_parse_sheet_csv_and_blank_row_skipping():
    csv_bytes = (
        "Team Name,Leader Name,Phone Number,Email\n"
        "Dragon Warriors,Praveen,9876543210,p@example.com\n"
        ",,,\n"
        "Border Runners,Asha,+91 91234 56780,asha@example.com\n"
    ).encode("utf-8")
    parsed = tis.parse_sheet(csv_bytes, "teams.csv")
    assert len(parsed.rows) == 2


def test_parse_sheet_rejects_legacy_xls():
    from app.core.exceptions import AppError

    with pytest.raises(AppError, match="re-save"):
        tis.parse_sheet(b"\xd0\xcf\x11\xe0", "teams.xls")


def test_parse_sheet_rejects_empty_upload():
    from app.core.exceptions import AppError

    with pytest.raises(AppError):
        tis.parse_sheet(b"", "teams.xlsx")


# ---------------------------------------------------------------------------
# Row building & validation
# ---------------------------------------------------------------------------


def _build(rows, *, existing_names=(), capacity=tis.MAX_TEAMS):
    data = _xlsx_bytes([["Team Name", "Leader Name", "Phone Number", "Email"], *rows])
    parsed = tis.parse_sheet(data, "teams.xlsx")
    mapping = tis.detect_columns(parsed.headers)
    return tis.build_rows(
        parsed,
        mapping,
        existing_names=existing_names,
        existing_codes=(),
        capacity=capacity,
    )


def test_build_rows_happy_path():
    rows = _build(
        [
            ["Dragon Warriors", "Praveen", "9876543210", "p@example.com"],
            ["Border Runners", "Asha", "+91 91234 56780", "asha@example.com"],
        ]
    )
    assert [r.status for r in rows] == [tis.STATUS_OK, tis.STATUS_OK]
    assert [r.password for r in rows] == ["9876", "9123"]
    assert rows[1].leader_phone == "9123456780"     # stored normalised
    assert all(r.team_code.startswith("B@GCEE-") for r in rows)
    assert rows[0].team_code != rows[1].team_code
    assert [r.row_number for r in rows] == [2, 3]


def test_build_rows_flags_missing_team_name():
    rows = _build([["", "Praveen", "9876543210", "p@example.com"]])
    assert rows[0].status == tis.STATUS_MISSING_TEAM_NAME


def test_build_rows_flags_unusable_phone():
    rows = _build([["Dragon Warriors", "Praveen", "12", "p@example.com"]])
    assert rows[0].status == tis.STATUS_INVALID_PHONE
    assert rows[0].team_code == ""      # no code burnt on a row that won't import


def test_build_rows_flags_duplicate_within_the_file():
    rows = _build(
        [
            ["Dragon Warriors", "A", "9876543210", ""],
            ["dragon  warriors", "B", "9123456780", ""],
        ]
    )
    assert rows[0].status == tis.STATUS_OK
    # Case-insensitive; the DB's UNIQUE(team_name) would reject it at commit.
    assert rows[1].status == tis.STATUS_DUPLICATE_IN_FILE


def test_build_rows_flags_duplicate_against_the_database():
    rows = _build(
        [["Dragon Warriors", "A", "9876543210", ""]],
        existing_names=["dragon warriors"],
    )
    assert rows[0].status == tis.STATUS_DUPLICATE_IN_DB


def test_build_rows_respects_remaining_capacity():
    rows = _build(
        [
            ["Team One", "A", "9876543210", ""],
            ["Team Two", "B", "9123456780", ""],
            ["Team Three", "C", "9111111111", ""],
        ],
        capacity=2,
    )
    assert [r.status for r in rows] == [
        tis.STATUS_OK,
        tis.STATUS_OK,
        tis.STATUS_OVER_CAPACITY,
    ]


def test_build_rows_never_reuses_a_code_across_rows():
    rows = _build([[f"Team {i}", "L", f"98765432{i:02d}", ""] for i in range(30)])
    codes = [r.team_code for r in rows if r.is_importable]
    assert len(codes) == len(set(codes)) == 30


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------


def _sheet_rows(content: bytes, sheet_name: str):
    workbook = load_workbook(io.BytesIO(content))
    return [list(r) for r in workbook[sheet_name].iter_rows(values_only=True)]


def _results_fixture():
    return [
        {
            "room_code": "R01", "room_number": 1, "rank": 1, "is_qualified": True,
            "team_code": "B@GCEE-1111#", "team_name": "Dragon Warriors",
            "leader_name": "Praveen", "leader_phone": "9876543210",
            "leader_email": "p@example.com", "suit_code": "HEART",
            "selected_number": 1, "mindmaze_score": 12.5,
            "king_diamond_score": 27.0, "jack_heart_score": 10.0,
            "total_score": 49.5,
        },
        {
            "room_code": "R01", "room_number": 1, "rank": 2, "is_qualified": False,
            "team_code": "B@GCEE-2222#", "team_name": "Border Runners",
            "leader_name": "Asha", "leader_phone": "9123456780",
            "leader_email": "a@example.com", "suit_code": "SPADE",
            "selected_number": 1, "mindmaze_score": 8.0,
            "king_diamond_score": 20.0, "jack_heart_score": 5.0,
            "total_score": 33.0,
        },
        {
            "room_code": "R02", "room_number": 2, "rank": 1, "is_qualified": True,
            "team_code": "B@GCEE-3333#", "team_name": "Night Owls",
            "leader_name": "Karthik", "leader_phone": "9000000001",
            "leader_email": "k@example.com", "suit_code": "CLUB",
            "selected_number": 2, "mindmaze_score": 20.0,
            "king_diamond_score": 30.0, "jack_heart_score": 15.0,
            "total_score": 65.0,
        },
    ]


def test_winners_workbook_lists_only_qualifiers_seeded_by_total():
    content = export_service.build_winners_workbook(
        round_name="Card Games",
        round_number=1,
        results=_results_fixture(),
        generated_at=datetime(2026, 9, 7, 12, 0, tzinfo=timezone.utc),
    )
    rows = _sheet_rows(content, "Round 2 Qualifiers")
    header_index = next(i for i, r in enumerate(rows) if r[0] == "Seed")
    body = [r for r in rows[header_index + 1 :] if r[0] is not None]

    assert len(body) == 2                       # the non-qualifier is excluded
    assert [r[2] for r in body] == ["Night Owls", "Dragon Warriors"]  # 65.0 before 49.5
    assert [r[0] for r in body] == [1, 2]       # re-seeded across rooms
    assert body[0][6] == "9000000001"           # leader phone carried through


def test_winners_workbook_all_teams_sheet_keeps_everyone():
    content = export_service.build_winners_workbook(
        round_name="Card Games",
        round_number=1,
        results=_results_fixture(),
        generated_at=datetime(2026, 9, 7, 12, 0, tzinfo=timezone.utc),
    )
    rows = _sheet_rows(content, "All Teams")
    header_index = next(i for i, r in enumerate(rows) if r[0] == "Room")
    body = [r for r in rows[header_index + 1 :] if r[0] is not None]
    assert len(body) == 3
    # openpyxl reads an empty cell back as None, not "".
    assert [r[2] or "" for r in body] == ["YES", "", "YES"]


def test_winners_workbook_with_no_qualifiers_still_builds():
    results = [dict(r, is_qualified=False) for r in _results_fixture()]
    content = export_service.build_winners_workbook(
        round_name="Card Games",
        round_number=1,
        results=results,
        generated_at=datetime(2026, 9, 7, 12, 0, tzinfo=timezone.utc),
    )
    rows = _sheet_rows(content, "Round 2 Qualifiers")
    assert any("0 team(s) qualified" in str(r[0]) for r in rows if r and r[0])


def test_credentials_workbook_round_trips():
    content = export_service.build_credentials_workbook(
        teams=[
            {
                "team_code": "B@GCEE-1111#", "password": "9876",
                "team_name": "Dragon Warriors", "leader_name": "Praveen",
                "leader_phone": "9876543210", "leader_email": "p@example.com",
            }
        ],
        generated_at=datetime(2026, 9, 7, 12, 0, tzinfo=timezone.utc),
    )
    rows = _sheet_rows(content, "Team Logins")
    header_index = next(i for i, r in enumerate(rows) if r[0] == "Team Code")
    assert rows[header_index + 1][:3] == ["B@GCEE-1111#", "9876", "Dragon Warriors"]


def test_import_template_is_readable_by_the_importer():
    """The template we hand organisers must survive its own round trip."""
    template = export_service.build_import_template()
    parsed = tis.parse_sheet(template, "template.xlsx")
    mapping = tis.detect_columns(parsed.headers)
    assert set(tis.REQUIRED_FIELDS) <= set(mapping)

    rows = tis.build_rows(
        parsed, mapping, existing_names=(), existing_codes=(), capacity=tis.MAX_TEAMS
    )
    sample = [r for r in rows if r.team_name]
    assert len(sample) == 2
    assert all(r.status == tis.STATUS_OK for r in sample)
    assert sample[0].password == "9876"
