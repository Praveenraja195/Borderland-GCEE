"""
Spreadsheet exports: the Round 2 qualifier list, and the credential sheet
handed out after a bulk team import.

Both produce a real .xlsx (openpyxl) rather than a CSV, because both are
documents a human forwards to other humans — a group chat, an event WhatsApp,
a co-organiser — and column widths, a frozen header and a readable title do
most of the work of making that forwardable.

The qualifier list is derived from `room_results`, which is the same table the
event's own qualification rule writes (`fn_compute_room_results`, applying
`qualification_rules.top_n`). Nothing here re-implements who qualified — it
reads the decision the event already made, optionally recomputing first so a
score corrected minutes ago is reflected.
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from typing import Any, Sequence

from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppError, NotFoundError
from app.models.round import Room, Round
from app.services import results_service

# Column widths are set by hand rather than measured: these sheets are opened
# on phones as often as on laptops, and a name column that fits the longest
# outlier makes every other column scroll off-screen.
_HEADER_FILL = "FF111827"
_HEADER_FONT = "FFFFFFFF"
_QUALIFIED_FILL = "FFDCFCE7"


def _openpyxl():
    try:
        import openpyxl  # noqa: F401
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise AppError(
            "Excel export is not available: openpyxl is not installed on the "
            "server. Run `pip install openpyxl` (already declared in "
            "pyproject.toml) and restart the API."
        ) from exc
    return openpyxl, Alignment, Font, PatternFill, get_column_letter


def _write_sheet(
    sheet,
    *,
    title_lines: Sequence[str],
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    widths: Sequence[int],
    highlight_rows: set[int] | None = None,
) -> None:
    """Lay one table out on a worksheet: title block, header, body, widths."""
    _openpyxl_mod, Alignment, Font, PatternFill, get_column_letter = _openpyxl()

    row_cursor = 1
    for index, line in enumerate(title_lines):
        cell = sheet.cell(row=row_cursor, column=1, value=line)
        cell.font = Font(bold=index == 0, size=14 if index == 0 else 10, color="FF334155")
        row_cursor += 1
    if title_lines:
        row_cursor += 1  # blank spacer row

    header_row = row_cursor
    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=header)
        cell.font = Font(bold=True, color=_HEADER_FONT)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row_cursor += 1

    highlight = PatternFill("solid", fgColor=_QUALIFIED_FILL)
    for body_index, row in enumerate(rows):
        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_cursor, column=col_index, value=value)
            if highlight_rows and body_index in highlight_rows:
                cell.fill = highlight
        row_cursor += 1

    for col_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col_index)].width = width

    # Freeze everything above and including the header so the table scrolls
    # under its own column names.
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if rows:
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
        )


def _workbook_bytes(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Round 2 qualifiers
# ---------------------------------------------------------------------------

_RESULTS_QUERY = text("""
    SELECT
        rm.room_code,
        rm.room_number,
        rr.rank,
        rr.is_qualified,
        t.team_code,
        t.team_name,
        t.leader_name,
        t.leader_phone,
        t.leader_email,
        su.code            AS suit_code,
        sel.selected_number,
        rr.mindmaze_score,
        rr.ace_spade_score,
        rr.king_diamond_score,
        rr.jack_heart_score,
        rr.total_score
    FROM room_results rr
    JOIN rooms rm ON rm.room_id = rr.room_id
    JOIN teams t  ON t.team_id  = rr.team_id
    LEFT JOIN round1_selections sel
           ON sel.team_id = t.team_id AND sel.round_id = rm.round_id
    LEFT JOIN suits su ON su.suit_id = sel.suit_id
    WHERE rm.round_id = :round_id
    ORDER BY rm.room_number, rr.rank NULLS LAST, t.team_code
""")


async def _fetch_round(db: AsyncSession, round_id: uuid.UUID) -> Round:
    round_obj = await db.get(Round, round_id)
    if round_obj is None:
        raise NotFoundError("Round not found")
    return round_obj


async def collect_round_results(
    db: AsyncSession, *, round_id: uuid.UUID, recompute: bool = True
) -> list[dict]:
    """Every team's final standing in a round, qualifiers flagged.

    `recompute` re-runs `fn_compute_room_results` for each room first. It
    defaults to on because the export is normally taken right after a score
    correction or a late force-close, and a stale `room_results` row would
    quietly export the wrong winner — the expensive-but-correct default is
    the right one for a document people act on.
    """
    await _fetch_round(db, round_id)

    if recompute:
        room_ids = (
            (await db.execute(select(Room.room_id).where(Room.round_id == round_id))).scalars().all()
        )
        for room_id in room_ids:
            await results_service.recompute_room_results(db, room_id=room_id)

    rows = (await db.execute(_RESULTS_QUERY, {"round_id": str(round_id)})).mappings().all()
    return [dict(row) for row in rows]


def _float(value: Any) -> float:
    return round(float(value or 0), 2)


def build_winners_workbook(
    *, round_name: str, round_number: int, results: Sequence[dict], generated_at: datetime
) -> bytes:
    """Two sheets: the qualifiers to carry into Round 2, and full standings.

    The qualifier sheet is seeded — sorted by total score across every room,
    not room by room — because that is the order Round 2 needs them in. The
    full-standings sheet is kept alongside it so any dispute ("why isn't my
    team on the list?") is answerable from the same file.
    """
    openpyxl, _Alignment, _Font, _PatternFill, _letter = _openpyxl()
    workbook = openpyxl.Workbook()

    stamp = generated_at.strftime("%d %b %Y, %H:%M UTC")
    qualifiers = [r for r in results if r.get("is_qualified")]
    qualifiers.sort(key=lambda r: (-_float(r.get("total_score")), str(r.get("team_code") or "")))

    winners_sheet = workbook.active
    winners_sheet.title = "Round 2 Qualifiers"
    _write_sheet(
        winners_sheet,
        title_lines=[
            f"Round {round_number} — {round_name}: Qualifiers for Round 2",
            f"{len(qualifiers)} team(s) qualified out of {len(results)}. Generated {stamp}.",
        ],
        headers=[
            "Seed",
            "Team Code",
            "Team Name",
            "Room",
            "Room Rank",
            "Leader",
            "Phone",
            "Email",
            "MindMaze",
            "Ace of Spades",
            "King of Diamonds",
            "Jack of Hearts",
            "Total",
        ],
        rows=[
            [
                seed,
                row.get("team_code"),
                row.get("team_name"),
                row.get("room_code"),
                row.get("rank"),
                row.get("leader_name"),
                row.get("leader_phone"),
                row.get("leader_email"),
                _float(row.get("mindmaze_score")),
                _float(row.get("ace_spade_score")),
                _float(row.get("king_diamond_score")),
                _float(row.get("jack_heart_score")),
                _float(row.get("total_score")),
            ]
            for seed, row in enumerate(qualifiers, start=1)
        ],
        widths=[6, 16, 26, 8, 11, 22, 15, 28, 11, 15, 17, 15, 10],
    )

    all_sheet = workbook.create_sheet("All Teams")
    qualified_indexes = {i for i, r in enumerate(results) if r.get("is_qualified")}
    _write_sheet(
        all_sheet,
        title_lines=[
            f"Round {round_number} — {round_name}: Full standings",
            f"All {len(results)} team(s), by room. Qualifiers are highlighted. Generated {stamp}.",
        ],
        headers=[
            "Room",
            "Rank",
            "Qualified",
            "Team Code",
            "Team Name",
            "Suit",
            "Number",
            "Leader",
            "Phone",
            "MindMaze",
            "Ace of Spades",
            "King of Diamonds",
            "Jack of Hearts",
            "Total",
        ],
        rows=[
            [
                row.get("room_code"),
                row.get("rank"),
                "YES" if row.get("is_qualified") else "",
                row.get("team_code"),
                row.get("team_name"),
                row.get("suit_code"),
                row.get("selected_number"),
                row.get("leader_name"),
                row.get("leader_phone"),
                _float(row.get("mindmaze_score")),
                _float(row.get("ace_spade_score")),
                _float(row.get("king_diamond_score")),
                _float(row.get("jack_heart_score")),
                _float(row.get("total_score")),
            ]
            for row in results
        ],
        widths=[8, 7, 10, 16, 26, 9, 9, 22, 15, 11, 15, 17, 15, 10],
        highlight_rows=qualified_indexes,
    )

    return _workbook_bytes(workbook)


async def build_winners_export(
    db: AsyncSession, *, round_id: uuid.UUID, recompute: bool = True
) -> tuple[bytes, str]:
    """The Round 2 qualifier workbook plus the filename to serve it under."""
    round_obj = await _fetch_round(db, round_id)
    results = await collect_round_results(db, round_id=round_id, recompute=recompute)
    generated_at = datetime.now(timezone.utc)
    content = build_winners_workbook(
        round_name=round_obj.name,
        round_number=round_obj.round_number,
        results=results,
        generated_at=generated_at,
    )
    filename = (
        f"round{round_obj.round_number}-qualifiers-"
        f"{generated_at.strftime('%Y%m%d-%H%M')}.xlsx"
    )
    return content, filename


# ---------------------------------------------------------------------------
# Team credentials (post-import handout)
# ---------------------------------------------------------------------------


def build_credentials_workbook(
    *, teams: Sequence[dict], generated_at: datetime, title: str = "Team Logins"
) -> bytes:
    """The sheet an organiser prints or forwards so each team can log in.

    Passwords are in the clear here by necessity — they are the first 4 digits
    of the leader's own phone number, which the leader already knows, and the
    hash in the database is one-way. Treat the file accordingly.
    """
    openpyxl, _Alignment, _Font, _PatternFill, _letter = _openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Team Logins"

    _write_sheet(
        sheet,
        title_lines=[
            title,
            f"{len(teams)} team(s). Password = first 4 digits of the leader's phone number. "
            f"Generated {generated_at.strftime('%d %b %Y, %H:%M UTC')}.",
        ],
        headers=[
            "Team Code",
            "Password",
            "Team Name",
            "Leader",
            "Phone",
            "Email",
        ],
        rows=[
            [
                team.get("team_code"),
                team.get("password"),
                team.get("team_name"),
                team.get("leader_name"),
                team.get("leader_phone"),
                team.get("leader_email"),
            ]
            for team in teams
        ],
        widths=[16, 12, 28, 22, 16, 30],
    )
    return _workbook_bytes(workbook)


def build_import_template() -> bytes:
    """A blank registration sheet in exactly the shape the importer expects.

    Handing the organiser this file removes the whole class of "the import
    couldn't find your phone column" problems before it starts.
    """
    openpyxl, Alignment, Font, _PatternFill, _letter = _openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Teams"

    _write_sheet(
        sheet,
        title_lines=[],
        headers=["Team Name", "Leader Name", "Phone Number", "Email"],
        rows=[
            ["Dragon Warriors", "Praveen Raja", "9876543210", "praveen@example.com"],
            ["Border Runners", "Asha Menon", "9123456780", "asha@example.com"],
        ],
        widths=[28, 24, 18, 30],
    )

    # The instructions live on their own sheet, not below the table. The
    # importer reads the *active* sheet top to bottom, so a note parked under
    # the sample rows would be parsed as a team named "Delete the two sample
    # rows..." — see test_import_template_is_readable_by_the_importer.
    help_sheet = workbook.create_sheet("How to use")
    help_sheet.column_dimensions["A"].width = 96
    for row_index, line in enumerate(
        [
            "Filling in this sheet",
            "",
            "1. Delete the two sample rows on the 'Teams' sheet, then add one row per team.",
            "2. Column order does not matter — columns are matched by their header names,",
            "   so you can also just upload your existing registration sheet as-is.",
            "3. Team Name and Phone Number are required. Leader Name and Email are optional.",
            "",
            "What the server fills in for you",
            "",
            "  Team Code — generated automatically, in the form B@GCEE-1234#",
            "  Password  — the first 4 digits of the leader's phone number",
            "              (e.g. 98765 43210 and +91 98765 43210 both give 9876)",
            "",
            "Nothing is created until you review the preview and confirm the import.",
        ],
        start=1,
    ):
        cell = help_sheet.cell(row=row_index, column=1, value=line)
        cell.alignment = Alignment(wrap_text=False, vertical="top")
        if line and not line.startswith((" ", "1", "2", "3")):
            cell.font = Font(bold=True, color="FF111827")
        else:
            cell.font = Font(color="FF334155")

    workbook.active = 0  # the importer must land on 'Teams'
    return _workbook_bytes(workbook)
